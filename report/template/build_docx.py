#!/usr/bin/env python3
"""Puts the report into the college template (Assignment 2_Report - Format - For Students.pdf).

Input : report/template/format.pdf   (the official 2-page format: cover page + table of contents)
        report/template/body.docx    (pandoc output of the report body, sections 1-9)
Output: report/BDA_Assignment2_Team16_Final.docx

Cover page and contents page reproduce the template exactly (same wording, red/black Times New Roman,
VTU and JSS logos taken from the PDF, page border). The "Page No." column holds PAGEREF fields that
point at the section headings, so Word fills in the real page numbers when the file is opened.
"""
import copy
import os

import fitz  # PyMuPDF
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

T = "report/template"
RED = RGBColor(0xC0, 0x00, 0x00)
FONT = "Times New Roman"

TITLE = ("Distributed Big Data Analytics of NOAA GHCN-Daily Climate Records using MongoDB, Hadoop MapReduce, "
         "Apache Pig and Apache Hive on a Multi-Node Cluster")
MEMBERS = [("Deeptanshu Pandey", "1JS23CS044"), ("Chirag S", "1JS23CS039"), ("Gagan Surendra Shamnur", "1JS23CS052")]
SECTIONS = ["Introduction and Objectives", "Dataset Description", "Distributed Computing Environment Setup",
            "Implementation of the Selected Practical Extension", "Analysis, Results, and Discussion",
            "Performance Evaluation (where applicable)", "Conclusion", "References",
            "Appendices – Code and Screenshots"]


# ------------------------------------------------------------------ logos from the template PDF
def extract_logos():
    d = fitz.open(f"{T}/format.pdf")
    page = d[0]
    imgs = []
    for img in page.get_images(full=True):
        bbox = page.get_image_bbox(img)
        pix = fitz.Pixmap(d, img[0])
        if pix.n - pix.alpha >= 4:
            pix = fitz.Pixmap(fitz.csRGB, pix)
        imgs.append((bbox.y0, bbox, pix))
    imgs.sort(key=lambda t: t[0])           # top image = VTU logo, bottom image = JSS logo
    vtu, jss = f"{T}/vtu_logo.png", f"{T}/jss_logo.png"
    imgs[0][2].save(vtu)
    imgs[-1][2].save(jss)
    print("logos:", [(round(b.x0), round(b.y0), round(b.width), round(b.height)) for _, b, _ in imgs])
    return vtu, jss


# ------------------------------------------------------------------ helpers
def para(doc, text="", size=12, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=4):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if text:
        r = p.add_run(text)
        r.font.name = FONT
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        r.font.size = Pt(size)
        r.bold = bold
        if color:
            r.font.color.rgb = color
    return p


def run(p, text, size=12, bold=False, color=None):
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = color
    return r


def cell_text(cell, text, bold=False, size=12, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    run(p, text, size=size, bold=bold)


def table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), "000000")
        borders.append(e)
    tblPr.append(borders)


def row_height(row, cm):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(cm * 567)))
    trPr.append(h)


def pageref(paragraph, bookmark):
    """PAGEREF field -> real page number of the bookmark (updated by Word on open)."""
    for kind, text in (("begin", None), (None, f" PAGEREF {bookmark} \\h "), ("separate", None), (None, "#"), ("end", None)):
        r = paragraph.add_run()
        r.font.name = FONT
        r.font.size = Pt(12)
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)
        elif text.startswith(" PAGEREF"):
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            r._r.append(it)
        else:
            r.text = text


def page_field(paragraph):
    for kind, text in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
        r = paragraph.add_run()
        r.font.name = FONT
        r.font.size = Pt(11)
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)
        elif text == " PAGE ":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            r._r.append(it)
        else:
            r.text = text


def page_border(sectPr):
    pb = OxmlElement("w:pgBorders")
    pb.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "8")
        e.set(qn("w:space"), "24")
        e.set(qn("w:color"), "000000")
        pb.append(e)
    sectPr.append(pb)


# ------------------------------------------------------------------ build
def heading_bookmarks(doc):
    """Bookmark each numbered section heading ("1. Introduction ...") as _Sec1.._Sec9 for PAGEREF."""
    names = [None] * len(SECTIONS)
    for p in doc.paragraphs:
        # the report's numbered sections ("## 1. Introduction ...") are Heading 2 in pandoc's output
        if not (p.style.name.lower().startswith("heading") and p.text[:2].rstrip(".").isdigit() and ". " in p.text[:4]):
            continue
        i = int(p.text.split(".")[0])
        if 1 <= i <= len(SECTIONS) and names[i - 1] is None:
            name = f"_Sec{i}"
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), str(900 + i))
            start.set(qn("w:name"), name)
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), str(900 + i))
            pPr = p._p.find(qn("w:pPr"))
            idx = list(p._p).index(pPr) + 1 if pPr is not None else 0
            p._p.insert(idx, start)
            p._p.append(end)
            names[i - 1] = name
    return names


def main():
    vtu, jss = extract_logos()
    doc = Document(f"{T}/body.docx")

    # ---- body styling: Times New Roman like the template
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        try:
            st = doc.styles[name]
            st.font.name = FONT
            st.font.size = Pt(12)
            st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        except KeyError:
            pass
    for lvl, size in ((1, 18), (2, 16), (3, 14), (4, 13)):
        try:
            st = doc.styles[f"Heading {lvl}"]
            st.font.name = FONT
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.color.rgb = RGBColor(0, 0, 0) if lvl > 2 else RED
            rf = st.element.rPr.rFonts
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(a)) is not None:
                    del rf.attrib[qn(a)]
            rf.set(qn("w:ascii"), FONT)
            rf.set(qn("w:hAnsi"), FONT)
        except KeyError:
            pass

    # pandoc writes no page size: A4 with the margins used throughout
    for sct in doc.sections:
        sct.page_width, sct.page_height = Cm(21.0), Cm(29.7)
        sct.left_margin = sct.right_margin = Cm(2.3)
        sct.top_margin = sct.bottom_margin = Cm(2.0)

    bookmarks = heading_bookmarks(doc)
    print("bookmarks:", bookmarks)

    body = doc.element.body
    original_elems = list(body)          # keep references alive so id() stays unique
    original = {id(e) for e in original_elems}

    # ================= COVER PAGE (template page 1) =================
    para(doc, "VISVESVARAYA TECHNOLOGICAL UNIVERSITY", 18, color=RED, after=0)
    para(doc, "Jnana Sangama, Santhibastawad Road, Machhe, Belagavi-590018, Karnataka, India", 11, after=6)
    p = para(doc, after=4)
    p.add_run().add_picture(vtu, width=Cm(2.6))
    para(doc, "Big Data Analytics", 14, bold=True, after=4)
    para(doc, "2022 SCHEME  -  BCS714D", 12, after=4)
    para(doc, "ASSIGNMENT 2 REPORT", 13, bold=True, after=4)
    para(doc, "ON", 11, after=4)
    para(doc, f"“{TITLE}”", 15, color=RED, after=6)
    para(doc, "Submitted as part of the coursework for Big Data Analytics", 12, after=0)
    para(doc, "BACHELOR OF ENGINEERING", 15, color=RED, after=0)
    para(doc, "in", 11, after=0)
    para(doc, "COMPUTER SCIENCE AND ENGINEERING", 15, color=RED, after=8)
    para(doc, "For the Academic Year 2026 - 2027", 12, after=6)
    p = para(doc, after=8)
    run(p, "Team NO: ", 15, color=RED)
    run(p, "16", 15, color=RED, bold=True)

    t = doc.add_table(rows=1 + 4, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(t)
    widths = (Cm(7.2), Cm(5.0), Cm(4.4))
    for i, h in enumerate(("Name", "USN", "Signature")):
        cell_text(t.rows[0].cells[i], h, bold=True)
    for r_i, row in enumerate(t.rows):
        row_height(row, 0.75 if r_i == 0 else 0.95)
        for c_i, c in enumerate(row.cells):
            c.width = widths[c_i]
    for r_i, (name, usn) in enumerate(MEMBERS, 1):
        cell_text(t.rows[r_i].cells[0], name)
        cell_text(t.rows[r_i].cells[1], usn)
        cell_text(t.rows[r_i].cells[2], "")

    p = para(doc, before=34, after=0, align=WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(12.6))
    run(p, "Signature of the Faculty:", 12, bold=True)
    run(p, "\tSignature of HOD:", 12, bold=True)
    para(doc, "Ms. Sandhiya B, Assistant Professor, Dept. of CSE", 10, align=WD_ALIGN_PARAGRAPH.LEFT, before=26, after=14)

    ft = doc.add_table(rows=1, cols=2)
    ft.alignment = WD_TABLE_ALIGNMENT.CENTER
    ft.rows[0].cells[0].width = Cm(2.2)
    ft.rows[0].cells[1].width = Cm(14.4)
    ft.rows[0].cells[0].paragraphs[0].add_run().add_picture(jss, width=Cm(1.6))
    c = ft.rows[0].cells[1]
    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(c.paragraphs[0], "Department of Computer Science and Engineering,\nJSS Academy of Technical Education,\n"
        "JSS Campus Rd, Dr. Vishnuvardhan Road, Bengaluru - 560060", 11)
    para(doc, after=0).add_run().add_break(WD_BREAK.PAGE)

    # ================= TABLE OF CONTENTS (template page 2) =================
    ht = doc.add_table(rows=1, cols=2)
    ht.alignment = WD_TABLE_ALIGNMENT.CENTER
    ht.rows[0].cells[0].width = Cm(2.2)
    ht.rows[0].cells[1].width = Cm(14.4)
    ht.rows[0].cells[0].paragraphs[0].add_run().add_picture(jss, width=Cm(1.6))
    c = ht.rows[0].cells[1]
    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(c.paragraphs[0], "JSS Academy of Technical Education\nDepartment of Computer Science and Engineering", 13)
    para(doc, "TABLE OF CONTENTS – ASSIGNMENT 2", 13, bold=True, before=22, after=12)

    toc = doc.add_table(rows=1 + len(SECTIONS), cols=3)
    toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(toc)
    tw = (Cm(2.6), Cm(11.0), Cm(3.0))
    for i, h in enumerate(("S. No.", "Section Title", "Page No.")):
        cell_text(toc.rows[0].cells[i], h, bold=True)
    for i, s in enumerate(SECTIONS, 1):
        cell_text(toc.rows[i].cells[0], str(i), bold=True)
        cell_text(toc.rows[i].cells[1], s)
        pc = toc.rows[i].cells[2]
        pc.text = ""
        pc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if bookmarks[i - 1]:
            pageref(pc.paragraphs[0], bookmarks[i - 1])
    for row in toc.rows:
        row_height(row, 0.9)
        for c_i, c in enumerate(row.cells):
            c.width = tw[c_i]

    # section break: cover + contents get the template's page border and no page number
    sp = doc.add_paragraph()
    sectPr = copy.deepcopy(doc.sections[-1]._sectPr)
    for child in list(sectPr):
        if child.tag in (qn("w:headerReference"), qn("w:footerReference"), qn("w:pgNumType")):
            sectPr.remove(child)
    page_border(sectPr)
    sp._p.get_or_add_pPr().append(sectPr)

    # move the new elements (everything appended after the original body) to the top
    # (python-docx inserts new blocks *before* the final sectPr, so select them by identity, not position)
    final_sectPr = body.find(qn("w:sectPr"))
    new = [e for e in body if id(e) not in original and e is not final_sectPr]
    for e in reversed(new):
        body.remove(e)
        body.insert(0, e)

    # body section: page numbers starting at 1, centered footer
    sec = doc.sections[-1]
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "1")
    sec._sectPr.append(pg)
    sec.footer.is_linked_to_previous = False
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_field(fp)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.3)
        s.top_margin = s.bottom_margin = Cm(2.0)

    # ask Word to refresh PAGEREF/PAGE fields when the document is opened
    settings = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    settings.append(uf)

    out = "report/BDA_Assignment2_Team16_Final.docx"
    doc.save(out)
    print("saved", out, os.path.getsize(out) // 1024, "KB")


if __name__ == "__main__":
    main()
